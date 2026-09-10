import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()
    
    # Page setup - Standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("LogSentinel: AntiGravity Execution Prompts")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("Full Implementation & Phased Prompts for AI-Driven Log Anomaly Detector")
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph() # Spacer

    def add_prompt_box(heading_text, prompt_text):
        h = doc.add_heading(heading_text, level=2)
        h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        
        # Add callout table/box for prompt content
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        
        cell = table.cell(0, 0)
        # Background color (Light Slate/Gray)
        shading = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        # Left border (Accent Blue)
        borders = parse_xml(
            r'<w:tcBorders {} >'
            r'<w:top w:val="none"/>'
            r'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
            r'<w:bottom w:val="none"/>'
            r'<w:right w:val="none"/>'
            r'</w:tcBorders>'.format(nsdecls('w'))
        )
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
        run = p.add_run(prompt_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
        doc.add_paragraph() # Spacer

    # Section 1: Backend + Frontend First Prompt
    p1_text = """Act as a Principal Full-Stack & Systems Engineer. Your task is to build and connect the complete FastAPI Backend and React Frontend for "LogSentinel — AI-Driven Log Anomaly Detector".

Create a fully working end-to-end web application with real-time WebSocket streaming, PostgreSQL persistence, and an interactive dark-mode dashboard.

File Structure to Implement:
- docker-compose.yml
- backend/
  - Dockerfile, requirements.txt
  - app/ (__init__.py, config.py, database.py, models.py, schemas.py, notifier.py, main.py)
  - scripts/ (simulate_logs.py)
- frontend/
  - Dockerfile, package.json, vite.config.js, tailwind.config.js, postcss.config.js, index.html
  - src/ (main.jsx, App.jsx, index.css)
  - src/components/ (MetricCards.jsx, RiskScoreChart.jsx, LiveLogFeed.jsx, AnomalyDetailModal.jsx, ComparisonPanel.jsx)

Requirements:
1. Backend (FastAPI + Async SQLAlchemy):
   - models: LogEntry (id, timestamp, raw_message, template_id, block_id, severity, is_anomaly, risk_score), AnomalyRecord (id, timestamp, risk_score, status, root_cause_chain, is_acknowledged).
   - endpoints: POST /api/logs/ingest, GET /api/anomalies, POST /api/anomalies/{id}/acknowledge, GET /api/metrics, ws://localhost:8000/ws/stream.
2. Frontend (React + Vite + Tailwind + Recharts):
   - LiveLogFeed (monospace scrolling feed with Drain3 tags), RiskScoreChart (0-100 line chart with threshold 75), MetricCards (logs/sec, risk status, anomaly count), AnomalyDetailModal (root-cause chain preview).
3. Simulator (backend/scripts/simulate_logs.py):
   - CLI tool supporting --rate and --mode (normal vs anomaly).
4. docker-compose.yml: Runs db (PostgreSQL 15), backend (port 8000), and frontend (port 3000) with volume mounts and auto-reload."""

    add_prompt_box("Option 1: Complete Backend + Frontend First Prompt", p1_text)

    # Section 2: Master Step-by-Step Prompt
    p2_text = """Act as a Principal Full-Stack, MLOps, and Distributed Systems Engineer. Build, validate, and containerize the complete production-grade application "LogSentinel — AI-Driven Log Anomaly Detector".

Step 1: Environment & Project Scaffolding
- Generate backend/, frontend/, dataset/, docker-compose.yml, and run_demo.sh.
- Configure PostgreSQL 15, FastAPI, PyTorch, Drain3, and Vite React frontend.

Step 2: Log Parsing & Data Pipeline
- Implement Drain3 TemplateMiner in backend/app/pipeline/parser.py with persistence.
- Implement sliding window and session/block grouping in backend/app/pipeline/windowing.py.

Step 3: PyTorch LSTM Autoencoder & Risk Scorer
- Implement LSTMAutoencoder (Embedding -> 2-layer LSTM -> Latent -> 2-layer LSTM -> Linear) in backend/app/ml/model.py.
- Implement training loop with early stopping, saving weights to backend/models/lstm_autoencoder.pt in backend/app/ml/train.py.
- Implement risk scoring (0-100) and root-cause chain extraction in backend/app/ml/scorer.py.

Step 4: Database Models & REST/WebSocket Endpoints
- Configure SQLAlchemy async models and endpoints in backend/app/main.py.
- Implement WebSocket broadcast on ws://localhost:8000/ws/stream.
- Implement asynchronous email alert dispatch with cooldown in backend/app/notifier.py.

Step 5: Dataset Download & Log Simulator
- Provide HDFS sample download script and autonomous log stream simulator in backend/scripts/simulate_logs.py.

Step 6: React Frontend Dashboard
- Build dark-mode interface in frontend/src/ with live streaming log terminal, Recharts risk score graph, metric tiles, root-cause explorer modal, and baseline comparison view.

Step 7: Automated Tests & Single-Command Demo
- Add integration unit tests in backend/tests/ and write run_demo.sh to orchestrate single-command startup."""

    add_prompt_box("Option 2: Master Step-by-Step System Prompt", p2_text)

    # Save file
    output_filename = "LogSentinel_AntiGravity_Prompts.docx"
    doc.save(output_filename)
    print(f"File successfully created: {output_filename}")

if __name__ == "__main__":
    create_document()
